import json
from pathlib import Path

from openpyxl import load_workbook

project_dir = Path(__file__).resolve().parent
wb_path = project_dir / 'json' / 'banco_de_dados.xlsx'
json_path = project_dir / 'public' / 'json' / 'banco_de_dados.json'

if not wb_path.exists():
    raise FileNotFoundError(f'Workbook not found: {wb_path}')

wb = load_workbook(wb_path, data_only=True)
records = {}

for sheet in wb.worksheets:
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item = {}
        for idx, value in enumerate(row):
            if idx < len(headers):
                item[headers[idx]] = value
        rows.append(item)
    records[sheet.title] = rows

json_path.parent.mkdir(parents=True, exist_ok=True)
json_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'created {json_path} with sheets: {list(records.keys())}')
