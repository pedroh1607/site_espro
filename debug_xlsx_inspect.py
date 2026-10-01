from openpyxl import load_workbook
import json

path = r"c:\Users\pedro.ribeiro\OneDrive - BGC Partners, O365 Tenant\OneDrive - BGCG, O365 Tenant\Documents\GitHub\site_espro\json\banco_de_dados.xlsx"
wb = load_workbook(path, data_only=True)
info = {}
for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    info[ws.title] = rows

out = r"c:\Users\pedro.ribeiro\OneDrive - BGC Partners, O365 Tenant\OneDrive - BGCG, O365 Tenant\Documents\GitHub\site_espro\debug_xlsx_inspect.json"
with open(out, 'w', encoding='utf-8') as f:
    json.dump(info, f, ensure_ascii=False, indent=2)
print(out)
